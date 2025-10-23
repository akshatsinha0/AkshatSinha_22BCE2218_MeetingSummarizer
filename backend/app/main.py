import os
import tempfile
import subprocess
import uuid
import json
import threading
import sqlite3
import datetime as dt
import warnings
import asyncio
import base64
from concurrent.futures import ThreadPoolExecutor
from typing import List, Optional

from fastapi import FastAPI, File, UploadFile, HTTPException, Form, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
from fastapi.staticfiles import StaticFiles

# Suppress warnings
warnings.filterwarnings("ignore", category=UserWarning)
warnings.filterwarnings("ignore", category=DeprecationWarning)
os.environ["CTRANSLATE2_VERBOSE"] = "0"

# Load .env from project root if present
try:
    here=os.path.abspath(os.path.dirname(__file__))
    root=os.path.abspath(os.path.join(here, os.pardir, os.pardir))
    load_dotenv(os.path.join(root, ".env"))
except Exception:
    pass

# ---- Config ----
OLLAMA_BASE_URL=os.getenv("OLLAMA_BASE_URL","http://127.0.0.1:11434")
OLLAMA_MODEL=os.getenv("OLLAMA_MODEL","qwen3:4b")
WHISPER_MODEL=os.getenv("WHISPER_MODEL","large-v3-turbo")
FFMPEG_BIN=os.getenv("FFMPEG_BIN","ffmpeg")
HF_TOKEN=os.getenv("HF_TOKEN","")
TRANSCRIBE_PROVIDER=os.getenv("TRANSCRIBE_PROVIDER","local") # local|azure|deepgram|google
DAILY_ASR_BUDGET_MIN=int(os.getenv("DAILY_ASR_BUDGET_MIN","0"))
CLOUD_FALLBACK=os.getenv("CLOUD_FALLBACK","never") # never|on_error|on_low_confidence
CONF_THRESHOLD=float(os.getenv("CONF_THRESHOLD","0.4"))
STORAGE_DIR=os.getenv("STORAGE_DIR",os.path.join(os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, os.pardir)),"storage"))
DB_PATH=os.path.join(STORAGE_DIR,"app.db")

os.makedirs(STORAGE_DIR,exist_ok=True)
os.makedirs(os.path.join(STORAGE_DIR,"audio"),exist_ok=True)
os.makedirs(os.path.join(STORAGE_DIR,"transcripts"),exist_ok=True)
os.makedirs(os.path.join(STORAGE_DIR,"summaries"),exist_ok=True)

app=FastAPI(title="Meeting Summarizer API")
app.mount("/storage", StaticFiles(directory=STORAGE_DIR), name="storage")
app.add_middleware(CORSMiddleware,allow_origins=["*"],allow_credentials=True,allow_methods=["*"],allow_headers=["*"])

class SummaryOut(BaseModel):
    transcript:str
    summary:str
    decisions:List[str]
    action_items:List[str]

class ReanalyzeRequest(BaseModel):
    transcript:str
    prompt:str

class JobOut(BaseModel):
    job_id:str
    status:str
    created_at:str
    updated_at:str
    input_path:Optional[str]=None
    transcript_path:Optional[str]=None
    summary_path:Optional[str]=None
    segments_path:Optional[str]=None
    model:Optional[str]=None
    language:Optional[str]=None
    diarization_enabled:Optional[bool]=None
    progress:Optional[float]=None
    stage:Optional[str]=None
    error:Optional[str]=None
    filename:Optional[str]=None

# ---- DB & Queue ----
_conn=sqlite3.connect(DB_PATH,check_same_thread=False)
_conn.row_factory=sqlite3.Row
_conn.execute("""
CREATE TABLE IF NOT EXISTS jobs (
  id TEXT PRIMARY KEY,
  status TEXT NOT NULL,
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL,
  input_path TEXT,
  transcript_path TEXT,
  summary_path TEXT,
  segments_path TEXT,
  model TEXT,
  language TEXT,
  diarization_enabled INTEGER,
  progress REAL,
  stage TEXT,
  error TEXT,
  filename TEXT
)
""")
# simple migration to add columns if missing
try:
    cols={r[1] for r in _conn.execute("PRAGMA table_info(jobs)").fetchall()}
    for name,type_ in [
        ("segments_path","TEXT"),("model","TEXT"),("language","TEXT"),("diarization_enabled","INTEGER"),("progress","REAL"),("stage","TEXT"),("filename","TEXT")
    ]:
        if name not in cols:
            _conn.execute(f"ALTER TABLE jobs ADD COLUMN {name} {type_}")
    _conn.commit()
except Exception:
    pass
_conn.execute("""
CREATE TABLE IF NOT EXISTS usage_budget (
  day TEXT PRIMARY KEY,
  minutes_used REAL NOT NULL
)
""")
_conn.execute("""
CREATE TABLE IF NOT EXISTS transcript_edits (
  id TEXT PRIMARY KEY,
  job_id TEXT NOT NULL,
  edited_transcript TEXT NOT NULL,
  created_at TEXT NOT NULL,
  FOREIGN KEY(job_id) REFERENCES jobs(id)
)
""")
_conn.execute("""
CREATE TABLE IF NOT EXISTS bookmarks (
  id TEXT PRIMARY KEY,
  job_id TEXT NOT NULL,
  timestamp REAL NOT NULL,
  label TEXT NOT NULL,
  created_at TEXT NOT NULL,
  FOREIGN KEY(job_id) REFERENCES jobs(id)
)
""")
_conn.commit()
_db_lock=threading.Lock()

_executor=ThreadPoolExecutor(max_workers=2)

# WebSocket connection manager for real-time transcription
class ConnectionManager:
    def __init__(self):
        self.active_connections: dict[str, WebSocket] = {}
    
    async def connect(self, session_id: str, websocket: WebSocket):
        await websocket.accept()
        self.active_connections[session_id] = websocket
    
    def disconnect(self, session_id: str):
        if session_id in self.active_connections:
            del self.active_connections[session_id]
    
    async def send_message(self, session_id: str, message: dict):
        if session_id in self.active_connections:
            try:
                await self.active_connections[session_id].send_json(message)
            except Exception:
                self.disconnect(session_id)

ws_manager = ConnectionManager()

# ---- Utilities ----

def _ensure_wav(input_path:str)->str:
    ext=os.path.splitext(input_path)[1].lower()
    if ext==".wav":
        return input_path
    wav_path=input_path+".wav"
    try:
        subprocess.run([FFMPEG_BIN,"-y","-i",input_path,"-ac","1","-ar","16000",wav_path],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    except FileNotFoundError:
        raise HTTPException(status_code=500,detail="ffmpeg not found. Set FFMPEG_BIN or ensure it is on PATH.")
    except subprocess.CalledProcessError as e:
        raise HTTPException(status_code=500,detail=f"ffmpeg failed: {e}")
    return wav_path

# Helpers for budget tracking

def _get_audio_duration(path:str)->float:
    try:
        import rust_audio
        duration, sample_rate, channels = rust_audio.validate_audio_file(path)
        return duration
    except Exception:
        try:
            r=subprocess.run([FFMPEG_BIN.replace('ffmpeg','ffprobe'),"-v","error","-show_entries","format=duration","-of","default=nw=1:nk=1",path],stdout=subprocess.PIPE,stderr=subprocess.PIPE,check=True,text=True)
            return float(r.stdout.strip())
        except Exception:
            return 0.0

def _budget_allow(duration_sec:float)->bool:
    if DAILY_ASR_BUDGET_MIN<=0:
        return False
    today=dt.date.today().isoformat()
    with _db_lock:
        cur=_conn.execute("SELECT minutes_used FROM usage_budget WHERE day=?",(today,))
        row=cur.fetchone()
        used=row[0] if row else 0.0
        allow=(used+duration_sec/60.0)<=DAILY_ASR_BUDGET_MIN
        return allow

def _budget_add(duration_sec:float)->None:
    today=dt.date.today().isoformat()
    with _db_lock:
        cur=_conn.execute("SELECT minutes_used FROM usage_budget WHERE day=?",(today,))
        row=cur.fetchone()
        used=row[0] if row else 0.0
        used+=duration_sec/60.0
        _conn.execute("REPLACE INTO usage_budget(day,minutes_used) VALUES(?,?)",(today,used))
        _conn.commit()

# ASR using faster-whisper (local)
def _transcribe(path:str,language=None):
    from faster_whisper import WhisperModel
    model=WhisperModel(WHISPER_MODEL,device="cpu")
    # Convert empty string to None for auto-detection
    lang = language if language and language.strip() else None
    segments,info=model.transcribe(path,beam_size=1,language=lang)
    out_segments=[]
    parts=[]
    for seg in segments:
        out_segments.append({"start":float(seg.start or 0),"end":float(seg.end or 0),"text":seg.text})
        parts.append(seg.text)
    return " ".join(parts).strip(),out_segments

# Cloud ASR (optional)

def _transcribe_cloud(path:str):
    provider=TRANSCRIBE_PROVIDER.lower()
    duration=_get_audio_duration(path)
    if provider=="local":
        return None
    if not _budget_allow(duration):
        return None
    try:
        import requests
        audio=open(path,'rb').read()
        if provider=="azure":
            key=os.getenv("AZURE_SPEECH_KEY"); region=os.getenv("AZURE_SPEECH_REGION")
            if not key or not region:
                return None
            url=f"https://{region}.stt.speech.microsoft.com/speech/recognition/conversation/cognitiveservices/v1?language=en-US"
            headers={"Ocp-Apim-Subscription-Key":key,"Content-Type":"audio/wav"}
            r=requests.post(url,data=audio,headers=headers,timeout=600)
            r.raise_for_status()
            _budget_add(duration)
            data=r.json(); text=data.get("DisplayText") or data.get("Text") or ""
            return text
        elif provider=="deepgram":
            dg=os.getenv("DEEPGRAM_API_KEY");
            if not dg:
                return None
            url="https://api.deepgram.com/v1/listen?model=nova-2-meeting&diarize=true&smart_format=true"
            headers={"Authorization":f"Token {dg}","Content-Type":"audio/wav"}
            r=requests.post(url,data=audio,headers=headers,timeout=600)
            r.raise_for_status()
            _budget_add(duration)
            data=r.json();
            text=" ".join(alt.get('transcript','') for alt in data.get('results',{}).get('channels',[{}])[0].get('alternatives',[]))
            return text
        elif provider=="google":
            try:
                from google.cloud import speech_v1p1beta1 as speech
                # GOOGLE_APPLICATION_CREDENTIALS should point to service account json
                client=speech.SpeechClient()
                audio_cfg=speech.RecognitionConfig(
                    encoding=speech.RecognitionConfig.AudioEncoding.LINEAR16,
                    sample_rate_hertz=16000,
                    language_code=os.getenv("GOOGLE_SPEECH_LANG","en-US"),
                    enable_automatic_punctuation=True,
                )
                audio_in=speech.RecognitionAudio(content=audio)
                resp=client.recognize(config=audio_cfg,audio=audio_in)
                _budget_add(duration)
                parts=[r.alternatives[0].transcript for r in resp.results if r.alternatives]
                return " ".join(parts)
            except Exception:
                return None
    except Exception:
        return None
    return None
# Summarization via Ollama HTTP API

def _summarize(transcript:str,*,model:Optional[str]=None,prompt_override:Optional[str]=None)->SummaryOut:
    import json,requests
    if prompt_override:
        prompt=f"{prompt_override}\n\nAnalyze the following transcript and return ONLY valid JSON with keys: summary, decisions (array), action_items (array).\n\nTranscript:\n{transcript}\n\nJSON:"
    else:
        prompt=(
            "You are a meeting summarizer. Analyze the transcript and provide:\n"
            "1. A concise summary of the meeting\n"
            "2. Key decisions made (if any)\n"
            "3. Action items with owners (if any)\n\n"
            "Return ONLY valid JSON in this exact format:\n"
            '{"summary": "text here", "decisions": ["decision 1", "decision 2"], "action_items": ["action 1", "action 2"]}\n\n'
            f"Transcript:\n{transcript}\n\nJSON:"
        )
    payload={"model":(model or OLLAMA_MODEL),"prompt":prompt,"stream":False,"format":"json"}
    try:
        r=requests.post(f"{OLLAMA_BASE_URL}/api/generate",json=payload,timeout=600)
        r.raise_for_status()
    except Exception as e:
        raise HTTPException(status_code=500,detail=f"Ollama request failed: {e}")
    data=r.json()
    txt=data.get("response","{}")
    parsed=None
    try:
        parsed=json.loads(txt)
    except Exception:
        import re
        m=re.search(r"\{[\s\S]*\}",txt)
        if m:
            try:
                parsed=json.loads(m.group(0))
            except Exception:
                pass
    if not isinstance(parsed,dict):
        parsed={"summary":txt[:800],"decisions":[],"action_items":[]}
    
    summary_text=str(parsed.get("summary",""))
    decisions_list=parsed.get("decisions",[])
    action_items_list=parsed.get("action_items",[])
    
    if not isinstance(decisions_list,list):
        decisions_list=[]
    if not isinstance(action_items_list,list):
        action_items_list=[]
    
    return SummaryOut(
        transcript=transcript,
        summary=summary_text,
        decisions=[str(d) for d in decisions_list],
        action_items=[str(a) for a in action_items_list]
    )

# Diarization via pyannote

def _diarize(path:str):
    if not HF_TOKEN:
        print("⚠️  HF_TOKEN not set - skipping diarization")
        return []
    try:
        from pyannote.audio import Pipeline
        print(f"<---->Loading diarization pipeline...")
        pipeline=Pipeline.from_pretrained("pyannote/speaker-diarization-3.1",use_auth_token=HF_TOKEN)
        print(f"<---->Running diarization on {path}...")
        annotation=pipeline(path)
        print(f"<---->Diarization complete")
    except Exception as e:
        # Fail soft if token missing or model access not granted
        print(f"<---->Diarization failed: {type(e).__name__}: {e}")
        return []
    diar_segments=[]
    speakers_map={}
    next_id=1
    for turn,_,speaker in annotation.itertracks(yield_label=True):
        if speaker not in speakers_map:
            speakers_map[speaker]=f"SPEAKER_{next_id}"
            next_id+=1
        diar_segments.append({"start":float(turn.start),"end":float(turn.end),"speaker":speakers_map[speaker]})
    diar_segments.sort(key=lambda x:x["start"])
    return diar_segments

def _assign_speakers(asr_segments,diar_segments):
    if not diar_segments:
        return [(None,s.get("text","")) for s in asr_segments]
    def find_speaker(ts):
        mid=(ts["start"]+ts["end"])/2.0
        for d in diar_segments:
            if d["start"]<=mid<=d["end"]:
                return d["speaker"]
        # fallback to max overlap
        best=None;over=0
        for d in diar_segments:
            a=max(d["start"],ts["start"]);b=min(d["end"],ts["end"])
            ov=max(0.0,b-a)
            if ov>over:
                over=ov;best=d["speaker"]
        return best
    labeled=[]
    for s in asr_segments:
        spk=find_speaker(s)
        labeled.append((spk,s["text"]))
    return labeled

@app.get("/health")
def health():
    return {"status":"ok","diarization":"enabled" if HF_TOKEN else "disabled","model":OLLAMA_MODEL}

# List available Ollama models
@app.get("/api/models")
def list_models():
    import requests
    try:
        r=requests.get(f"{OLLAMA_BASE_URL}/api/tags",timeout=30)
        r.raise_for_status()
        data=r.json()
        return data
    except Exception as e:
        raise HTTPException(status_code=500,detail=f"Unable to list models: {e}")

# Exports
@app.get("/api/jobs/{job_id}/export/pdf")
def export_pdf(job_id:str):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas
    with _db_lock:
        row=_conn.execute("SELECT summary_path,transcript_path FROM jobs WHERE id=?",(job_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404,detail="job not found")
    # read summary json
    try:
        import requests
        s=requests.get(f"http://127.0.0.1:8000{row[0]}").json()
    except Exception:
        s={}
    pdf_file=os.path.join(STORAGE_DIR,"summaries",f"{job_id}.pdf")
    c=canvas.Canvas(pdf_file,pagesize=LETTER)
    y=750
    def line(txt):
        nonlocal y
        c.drawString(40,y,txt[:100])
        y-=18
        if y<60:
            c.showPage(); y=750
    line("Meeting Summary")
    for l in (s.get('summary','') or '').split('\n'):
        line(l)
    line("")
    line("Decisions:")
    for d in s.get('decisions',[]) or []:
        line(f"- {d}")
    line("")
    line("Action Items:")
    for a in s.get('action_items',[]) or []:
        line(f"- {a}")
    c.save()
    return {"pdf":"/storage/summaries/%s.pdf"%job_id}

@app.get("/api/jobs/{job_id}/export/docx")
def export_docx(job_id:str):
    from docx import Document
    from docx.shared import Pt
    import requests
    with _db_lock:
        row=_conn.execute("SELECT summary_path FROM jobs WHERE id=?",(job_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404,detail="job not found")
    s=requests.get(f"http://127.0.0.1:8000{row[0]}").json()
    doc=Document()
    doc.add_heading('Meeting Summary', level=1)
    doc.add_paragraph(s.get('summary','') or '')
    doc.add_heading('Decisions', level=2)
    for d in s.get('decisions',[]) or []:
        doc.add_paragraph(d, style='List Bullet')
    doc.add_heading('Action Items', level=2)
    for a in s.get('action_items',[]) or []:
        doc.add_paragraph(a, style='List Bullet')
    docx_file=os.path.join(STORAGE_DIR,"summaries",f"{job_id}.docx")
    doc.save(docx_file)
    return {"docx":"/storage/summaries/%s.docx"%job_id}

@app.post("/api/process",response_model=SummaryOut)
async def process_meeting(
    file:UploadFile=File(...),
    model:Optional[str]=Form(None),
    language:Optional[str]=Form(None),
    diarization_enabled:Optional[bool]=Form(True),
    prompt:Optional[str]=Form(None)
):
    suffix=os.path.splitext(file.filename or "audio")[1] or ".wav"
    with tempfile.NamedTemporaryFile(delete=False,suffix=suffix) as tmp:
        raw_path=tmp.name
        content=await file.read()
        tmp.write(content)
    wav_path=_ensure_wav(raw_path)
    # try cloud only if allowed
    cloud_text=_transcribe_cloud(wav_path)
    if cloud_text:
        transcript=cloud_text
        segments=[{"start":0.0,"end":0.0,"text":cloud_text}]
    else:
        transcript,segments=_transcribe(wav_path,language=language)
    diar=_diarize(wav_path) if diarization_enabled else []
    labeled=_assign_speakers(segments,diar)
    labeled_text=[]
    for spk,txt in labeled:
        if spk: labeled_text.append(f"{spk}: {txt}")
        else: labeled_text.append(txt)
    final_transcript="\n".join(labeled_text) if labeled_text else transcript
    return _summarize(final_transcript,model=model,prompt_override=prompt)

# ---- Async Jobs API ----

def _db_upsert_job(job_id:str, **kwargs):
    now=dt.datetime.utcnow().isoformat()
    with _db_lock:
        cur=_conn.execute("SELECT id FROM jobs WHERE id=?",(job_id,))
        exists=cur.fetchone() is not None
        if exists:
            sets=", ".join([f"{k}=?" for k in kwargs.keys()])
            vals=list(kwargs.values())
            vals.append(job_id)
            _conn.execute(f"UPDATE jobs SET {sets}, updated_at=? WHERE id=?",(*kwargs.values(),now,job_id))
        else:
            _conn.execute("INSERT INTO jobs(id,status,created_at,updated_at) VALUES(?,?,?,?)",(job_id,kwargs.get('status','queued'),now,now))
        _conn.commit()

def _job_row_to_out(row)->JobOut:
    input_url = None
    if row['input_path']:
        filename = os.path.basename(row['input_path'])
        input_url = f"/storage/audio/{filename}"
    
    return JobOut(
        job_id=row['id'],
        status=row['status'],
        created_at=row['created_at'],
        updated_at=row['updated_at'],
        input_path=input_url,
        transcript_path=row['transcript_path'],
        summary_path=row['summary_path'],
        segments_path=row['segments_path'],
        model=row['model'],
        language=row['language'],
        diarization_enabled=bool(row['diarization_enabled']) if row['diarization_enabled'] is not None else None,
        progress=row['progress'],
        stage=row['stage'],
        error=row['error'],
        filename=row['filename'] if 'filename' in row.keys() else None
    )

def _process_job(job_id:str, file_path:str, *, model_override:Optional[str]=None, language:Optional[str]=None, diarization_enabled:bool=True, prompt_override:Optional[str]=None):
    try:
        _db_upsert_job(job_id,status='processing',input_path=file_path,stage='preprocess',progress=0.05)
        wav_path=_ensure_wav(file_path)
        _db_upsert_job(job_id,stage='transcribing',progress=0.25)
        print(f"Job {job_id}: Starting transcription...")
        cloud_text=_transcribe_cloud(wav_path)
        if cloud_text:
            transcript=cloud_text
            segments=[{"start":0.0,"end":0.0,"text":cloud_text}]
        else:
            transcript,segments=_transcribe(wav_path,language=language)
        print(f"Job {job_id}: Transcription complete")
        _db_upsert_job(job_id,stage='diarizing',progress=0.6)
        diar=_diarize(wav_path) if diarization_enabled else []
        labeled=_assign_speakers(segments,diar)
        labeled_text=[(f"{spk}: {txt}" if spk else txt) for spk,txt in labeled]
        final_transcript="\n".join(labeled_text) if labeled_text else transcript
        _db_upsert_job(job_id,stage='summarizing',progress=0.8)
        summary=_summarize(final_transcript,model=model_override,prompt_override=prompt_override)
        # persist
        segfile=os.path.join(STORAGE_DIR,"segments",f"{job_id}.json")
        os.makedirs(os.path.dirname(segfile),exist_ok=True)
        open(segfile,'w',encoding='utf-8').write(json.dumps({"segments":labeled, "raw":segments},indent=2))
        tfile=os.path.join(STORAGE_DIR,"transcripts",f"{job_id}.txt")
        sfile=os.path.join(STORAGE_DIR,"summaries",f"{job_id}.json")
        open(tfile,'w',encoding='utf-8').write(final_transcript)
        open(sfile,'w',encoding='utf-8').write(summary.model_dump_json(indent=2))
        turl=f"/storage/transcripts/{job_id}.txt"
        surl=f"/storage/summaries/{job_id}.json"
        segurl=f"/storage/segments/{job_id}.json"
        _db_upsert_job(job_id,status='done',transcript_path=turl,summary_path=surl,segments_path=segurl,model=model_override or OLLAMA_MODEL,language=language,diarization_enabled=1 if diarization_enabled else 0,stage='done',progress=1.0)
    except Exception as e:
        import traceback
        error_msg = f"{type(e).__name__}: {e}"
        print(f"Job {job_id} failed: {error_msg}")
        traceback.print_exc()
        _db_upsert_job(job_id,status='error',error=error_msg,stage='error')

@app.post("/api/jobs",response_model=JobOut)
async def create_job(
    file:UploadFile=File(...),
    model:Optional[str]=Form(None),
    language:Optional[str]=Form(None),
    diarization_enabled:Optional[bool]=Form(True),
    prompt:Optional[str]=Form(None)
):
    job_id=str(uuid.uuid4())
    original_filename=file.filename or "audio"
    suffix=os.path.splitext(original_filename)[1] or ".wav"
    audio_path=os.path.join(STORAGE_DIR,"audio",f"{job_id}{suffix}")
    with open(audio_path,'wb') as f:
        f.write(await file.read())
    _db_upsert_job(job_id,status='queued',input_path=audio_path,model=model or OLLAMA_MODEL,language=language,diarization_enabled=1 if diarization_enabled else 0,progress=0.0,stage='queued',filename=original_filename)
    _executor.submit(_process_job,job_id,audio_path,model_override=model,language=language,diarization_enabled=bool(diarization_enabled),prompt_override=prompt)
    with _db_lock:
        row=_conn.execute("SELECT * FROM jobs WHERE id=?",(job_id,)).fetchone()
    return _job_row_to_out(row)

@app.get("/api/jobs/{job_id}",response_model=JobOut)
async def get_job(job_id:str):
    with _db_lock:
        row=_conn.execute("SELECT * FROM jobs WHERE id=?",(job_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404,detail="job not found")
    return _job_row_to_out(row)

@app.get("/api/jobs",response_model=List[JobOut])
async def list_jobs(limit:int=20):
    with _db_lock:
        rows=_conn.execute("SELECT * FROM jobs ORDER BY created_at DESC LIMIT ?",(limit,)).fetchall()
    return [_job_row_to_out(r) for r in rows]

@app.post("/api/reanalyze")
async def reanalyze_transcript(req:ReanalyzeRequest):
    return _summarize(req.transcript,prompt_override=req.prompt)

# Real-time transcription WebSocket endpoint
@app.websocket("/ws/transcribe/{session_id}")
async def websocket_transcribe(websocket: WebSocket, session_id: str):
    await ws_manager.connect(session_id, websocket)
    audio_buffer = b""
    temp_files = []
    
    try:
        from faster_whisper import WhisperModel
        model = WhisperModel(WHISPER_MODEL, device="cpu")
        
        await ws_manager.send_message(session_id, {"type": "status", "message": "Connected. Start streaming audio..."})
        
        while True:
            data = await websocket.receive()
            
            if "text" in data:
                msg = json.loads(data["text"])
                if msg.get("type") == "audio_chunk":
                    # Decode base64 audio chunk
                    chunk = base64.b64decode(msg["data"])
                    audio_buffer += chunk
                    
                    # Process when buffer reaches ~3 seconds (48000 samples * 2 bytes)
                    if len(audio_buffer) >= 96000:
                        # Save to temp file
                        temp_path = os.path.join(tempfile.gettempdir(), f"{session_id}_{uuid.uuid4()}.wav")
                        temp_files.append(temp_path)
                        
                        with open(temp_path, "wb") as f:
                            f.write(audio_buffer)
                        
                        # Transcribe chunk
                        try:
                            segments, _ = model.transcribe(temp_path, beam_size=1)
                            text_parts = []
                            for seg in segments:
                                text_parts.append(seg.text)
                            
                            if text_parts:
                                await ws_manager.send_message(session_id, {
                                    "type": "transcript",
                                    "text": " ".join(text_parts).strip(),
                                    "timestamp": dt.datetime.utcnow().isoformat()
                                })
                        except Exception as e:
                            await ws_manager.send_message(session_id, {
                                "type": "error",
                                "message": f"Transcription error: {str(e)}"
                            })
                        
                        # Clear buffer
                        audio_buffer = b""
                
                elif msg.get("type") == "end":
                    await ws_manager.send_message(session_id, {"type": "status", "message": "Session ended"})
                    break
            
            elif "bytes" in data:
                # Direct binary audio data
                audio_buffer += data["bytes"]
                
                if len(audio_buffer) >= 96000:
                    temp_path = os.path.join(tempfile.gettempdir(), f"{session_id}_{uuid.uuid4()}.wav")
                    temp_files.append(temp_path)
                    
                    with open(temp_path, "wb") as f:
                        f.write(audio_buffer)
                    
                    try:
                        segments, _ = model.transcribe(temp_path, beam_size=1)
                        text_parts = [seg.text for seg in segments]
                        
                        if text_parts:
                            await ws_manager.send_message(session_id, {
                                "type": "transcript",
                                "text": " ".join(text_parts).strip(),
                                "timestamp": dt.datetime.utcnow().isoformat()
                            })
                    except Exception as e:
                        await ws_manager.send_message(session_id, {
                            "type": "error",
                            "message": f"Transcription error: {str(e)}"
                        })
                    
                    audio_buffer = b""
    
    except WebSocketDisconnect:
        ws_manager.disconnect(session_id)
    except Exception as e:
        await ws_manager.send_message(session_id, {"type": "error", "message": str(e)})
        ws_manager.disconnect(session_id)
    finally:
        # Cleanup temp files
        for temp_file in temp_files:
            try:
                if os.path.exists(temp_file):
                    os.remove(temp_file)
            except Exception:
                pass

# Transcript editing endpoints
class TranscriptEditRequest(BaseModel):
    job_id: str
    edited_transcript: str

@app.post("/api/jobs/{job_id}/edit-transcript")
async def edit_transcript(job_id: str, req: TranscriptEditRequest):
    edit_id = str(uuid.uuid4())
    now = dt.datetime.utcnow().isoformat()
    
    with _db_lock:
        # Save edit history
        _conn.execute(
            "INSERT INTO transcript_edits(id, job_id, edited_transcript, created_at) VALUES(?,?,?,?)",
            (edit_id, job_id, req.edited_transcript, now)
        )
        _conn.commit()
        
        # Update transcript file
        row = _conn.execute("SELECT transcript_path FROM jobs WHERE id=?", (job_id,)).fetchone()
        if row and row[0]:
            transcript_file = os.path.join(STORAGE_DIR, "transcripts", f"{job_id}.txt")
            with open(transcript_file, 'w', encoding='utf-8') as f:
                f.write(req.edited_transcript)
    
    return {"status": "success", "edit_id": edit_id}

@app.get("/api/jobs/{job_id}/edit-history")
async def get_edit_history(job_id: str):
    with _db_lock:
        rows = _conn.execute(
            "SELECT id, created_at FROM transcript_edits WHERE job_id=? ORDER BY created_at DESC",
            (job_id,)
        ).fetchall()
    
    return [{"edit_id": r[0], "created_at": r[1]} for r in rows]

@app.post("/api/jobs/{job_id}/rerun-summary")
async def rerun_summary_after_edit(job_id: str):
    with _db_lock:
        row = _conn.execute("SELECT transcript_path FROM jobs WHERE id=?", (job_id,)).fetchone()
    
    if not row or not row[0]:
        raise HTTPException(status_code=404, detail="Job not found")
    
    # Read edited transcript
    transcript_file = os.path.join(STORAGE_DIR, "transcripts", f"{job_id}.txt")
    with open(transcript_file, 'r', encoding='utf-8') as f:
        transcript = f.read()
    
    # Re-run summarization
    summary = _summarize(transcript)
    
    # Save new summary
    sfile = os.path.join(STORAGE_DIR, "summaries", f"{job_id}.json")
    with open(sfile, 'w', encoding='utf-8') as f:
        f.write(summary.model_dump_json(indent=2))
    
    return {"status": "success", "summary": summary}

# Bookmarks endpoints
class BookmarkRequest(BaseModel):
    timestamp: float
    label: str

@app.post("/api/jobs/{job_id}/bookmarks")
async def add_bookmark(job_id: str, req: BookmarkRequest):
    bookmark_id = str(uuid.uuid4())
    now = dt.datetime.utcnow().isoformat()
    
    with _db_lock:
        _conn.execute(
            "INSERT INTO bookmarks(id, job_id, timestamp, label, created_at) VALUES(?,?,?,?,?)",
            (bookmark_id, job_id, req.timestamp, req.label, now)
        )
        _conn.commit()
    
    return {"status": "success", "bookmark_id": bookmark_id}

@app.get("/api/jobs/{job_id}/bookmarks")
async def get_bookmarks(job_id: str):
    with _db_lock:
        rows = _conn.execute(
            "SELECT id, timestamp, label, created_at FROM bookmarks WHERE job_id=? ORDER BY timestamp",
            (job_id,)
        ).fetchall()
    
    return [{"id": r[0], "timestamp": r[1], "label": r[2], "created_at": r[3]} for r in rows]

@app.delete("/api/jobs/{job_id}/bookmarks/{bookmark_id}")
async def delete_bookmark(job_id: str, bookmark_id: str):
    with _db_lock:
        _conn.execute("DELETE FROM bookmarks WHERE id=? AND job_id=?", (bookmark_id, job_id))
        _conn.commit()
    
    return {"status": "success"}

# Batch processing endpoint
@app.post("/api/jobs/batch")
async def create_batch_jobs(
    files: List[UploadFile] = File(...),
    model: Optional[str] = Form(None),
    language: Optional[str] = Form(None),
    diarization_enabled: Optional[bool] = Form(True)
):
    job_ids = []
    
    for file in files:
        job_id = str(uuid.uuid4())
        original_filename = file.filename or "audio"
        suffix = os.path.splitext(original_filename)[1] or ".wav"
        audio_path = os.path.join(STORAGE_DIR, "audio", f"{job_id}{suffix}")
        
        with open(audio_path, 'wb') as f:
            f.write(await file.read())
        
        _db_upsert_job(
            job_id,
            status='queued',
            input_path=audio_path,
            model=model or OLLAMA_MODEL,
            language=language,
            diarization_enabled=1 if diarization_enabled else 0,
            progress=0.0,
            stage='queued',
            filename=original_filename
        )
        
        _executor.submit(
            _process_job,
            job_id,
            audio_path,
            model_override=model,
            language=language,
            diarization_enabled=bool(diarization_enabled),
            prompt_override=None
        )
        
        job_ids.append(job_id)
    
    return {"status": "success", "job_ids": job_ids, "count": len(job_ids)}

# Bulk export endpoint
@app.post("/api/jobs/bulk-export")
async def bulk_export(job_ids: List[str], format: str = "pdf"):
    results = []
    
    for job_id in job_ids:
        try:
            if format == "pdf":
                result = await export_pdf(job_id)
            elif format == "docx":
                result = await export_docx(job_id)
            else:
                continue
            
            results.append({"job_id": job_id, "url": result.get(format)})
        except Exception as e:
            results.append({"job_id": job_id, "error": str(e)})
    
    return {"results": results}

# Compliance checking endpoint
# Calendar integration endpoints
@app.get("/api/calendar/upcoming")
async def get_upcoming_meetings(max_results: int = 10):
    try:
        from .calendar_integration import get_calendar_service
        cal = get_calendar_service()
        meetings = cal.get_upcoming_meetings(max_results)
        return {"meetings": meetings}
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Calendar error: {str(e)}")

@app.get("/api/calendar/search")
async def search_meetings(query: str, max_results: int = 10):
    try:
        from .calendar_integration import get_calendar_service
        cal = get_calendar_service()
        meetings = cal.search_meetings_by_title(query, max_results)
        return {"meetings": meetings}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Calendar error: {str(e)}")

@app.post("/api/calendar/link-meeting")
async def link_meeting_to_job(job_id: str, calendar_event_id: str):
    """Link a calendar event to a job"""
    try:
        from .calendar_integration import get_calendar_service
        cal = get_calendar_service()
        meeting = cal.get_meeting_by_id(calendar_event_id)
        
        if not meeting:
            raise HTTPException(status_code=404, detail="Calendar event not found")
        
        # Store meeting metadata with job
        with _db_lock:
            # Add calendar_event_id column if not exists
            try:
                _conn.execute("ALTER TABLE jobs ADD COLUMN calendar_event_id TEXT")
                _conn.execute("ALTER TABLE jobs ADD COLUMN meeting_metadata TEXT")
                _conn.commit()
            except Exception:
                pass
            
            _conn.execute(
                "UPDATE jobs SET calendar_event_id=?, meeting_metadata=? WHERE id=?",
                (calendar_event_id, json.dumps(meeting), job_id)
            )
            _conn.commit()
        
        return {"status": "success", "meeting": meeting}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error linking meeting: {str(e)}")

@app.get("/api/jobs/{job_id}/compliance-check")
async def compliance_check(job_id: str):
    with _db_lock:
        row = _conn.execute("SELECT transcript_path FROM jobs WHERE id=?", (job_id,)).fetchone()
    
    if not row or not row[0]:
        raise HTTPException(status_code=404, detail="Job not found")
    
    # Read transcript
    transcript_file = os.path.join(STORAGE_DIR, "transcripts", f"{job_id}.txt")
    with open(transcript_file, 'r', encoding='utf-8') as f:
        transcript = f.read()
    
    # Compliance patterns
    issues = []
    
    # PII detection patterns
    import re
    
    # Email addresses
    emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', transcript)
    if emails:
        issues.append({
            "type": "PII",
            "severity": "high",
            "description": f"Found {len(emails)} email address(es)",
            "examples": emails[:3]
        })
    
    # Phone numbers
    phones = re.findall(r'\b\d{3}[-.]?\d{3}[-.]?\d{4}\b', transcript)
    if phones:
        issues.append({
            "type": "PII",
            "severity": "high",
            "description": f"Found {len(phones)} phone number(s)",
            "examples": phones[:3]
        })
    
    # SSN patterns
    ssns = re.findall(r'\b\d{3}-\d{2}-\d{4}\b', transcript)
    if ssns:
        issues.append({
            "type": "PII",
            "severity": "critical",
            "description": f"Found {len(ssns)} potential SSN(s)",
            "examples": ["***-**-****"] * len(ssns[:3])
        })
    
    # Legal keywords
    legal_keywords = ['confidential', 'proprietary', 'trade secret', 'NDA', 'non-disclosure']
    found_legal = [kw for kw in legal_keywords if kw.lower() in transcript.lower()]
    if found_legal:
        issues.append({
            "type": "Legal",
            "severity": "medium",
            "description": "Contains legal/confidential terms",
            "examples": found_legal
        })
    
    # Financial keywords
    financial_keywords = ['salary', 'compensation', 'budget', 'revenue', 'profit', 'loss']
    found_financial = [kw for kw in financial_keywords if kw.lower() in transcript.lower()]
    if found_financial:
        issues.append({
            "type": "Financial",
            "severity": "medium",
            "description": "Contains financial information",
            "examples": found_financial
        })
    
    return {
        "job_id": job_id,
        "issues": issues,
        "total_issues": len(issues),
        "status": "clean" if len(issues) == 0 else "flagged"
    }
